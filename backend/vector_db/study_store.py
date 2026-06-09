# from sentence_transformers import SentenceTransformer
# import faiss
# import numpy as np

# model = SentenceTransformer("all-MiniLM-L6-v2")

# store = {}  # {(user_id, subject): {docs, index}}

# def add_content(user_id, subject, text):
#     key = (user_id, subject)

#     if key not in store:
#         store[key] = {
#             "documents": [],
#             "index": faiss.IndexFlatL2(384)
#         }

#     emb = model.encode([text])
#     store[key]["documents"].append(text)
#     store[key]["index"].add(emb)

# def retrieve_context(user_id, subject, k=3):
#     key = (user_id, subject)
#     if key not in store:
#         return ""

#     docs = store[key]["documents"]
#     index = store[key]["index"]

#     q_emb = model.encode([subject])
#     _, idx = index.search(q_emb, k)

#     return "\n".join([docs[i] for i in idx[0]])

























# from sentence_transformers import SentenceTransformer
# import faiss
# import numpy as np
# import PyPDF2
# import docx

# model = SentenceTransformer("all-MiniLM-L6-v2")

# # Store format: {(user_id, subject): {"documents": [...], "index": FAISS index}}
# store = {}

# def add_content(user_id, subject, text):
#     key = (user_id, subject)
#     if key not in store:
#         store[key] = {
#             "documents": [],
#             "index": faiss.IndexFlatL2(384)
#         }

#     emb = model.encode([text])
#     store[key]["documents"].append(text)
#     store[key]["index"].add(np.array(emb).astype("float32"))

# def retrieve_context(user_id, subject, k=3):
#     key = (user_id, subject)
#     if key not in store or len(store[key]["documents"]) == 0:
#         return ""

#     docs = store[key]["documents"]
#     index = store[key]["index"]

#     q_emb = model.encode([subject])
#     _, idx = index.search(np.array(q_emb).astype("float32"), k)
#     return "\n".join([docs[i] for i in idx[0]])

# def add_content_from_file(user_id, subject, file):
#     """Support PDF, DOCX, TXT"""
#     text_to_add = ""

#     if file.filename.endswith(".txt"):
#         text_to_add = file.file.read().decode("utf-8")

#     elif file.filename.endswith(".pdf"):
#         reader = PyPDF2.PdfReader(file.file)
#         text_to_add = " ".join(
#             [page.extract_text() or "" for page in reader.pages]
#         )

#     elif file.filename.endswith(".docx"):
#         doc = docx.Document(file.file)
#         text_to_add = " ".join([p.text for p in doc.paragraphs])

#     else:
#         raise ValueError("Unsupported file type")

#     # ✅ SAFETY CHECK (ADDITION)
#     text_to_add = text_to_add.strip()
#     if not text_to_add:
#         return ""

#     add_content(user_id, subject, text_to_add)

#     # ✅ RETURN TEXT (CRITICAL ADDITION)
#     return text_to_add





import os
import faiss
import numpy as np
from sentence_transformers import SentenceTransformer
from fastapi import UploadFile
import pdfplumber
from docx import Document

# -----------------------------
# BASE SETUP
# -----------------------------
BASE_DIR = "backend/vector_db/data"
os.makedirs(BASE_DIR, exist_ok=True)

model = SentenceTransformer("all-MiniLM-L6-v2")

# IMPORTANT: now structured per user + subject
indexes = {}
documents = {}

# -----------------------------
# PATH (future use if you want persistence)
# -----------------------------
def _get_user_path(user_id: str):
    path = os.path.join(BASE_DIR, str(user_id))
    os.makedirs(path, exist_ok=True)
    return path


# -----------------------------
# ADD CONTENT (FIXED)
# -----------------------------
def add_content(user_id: str, subject: str, text: str):

    if not text or not text.strip():
        return

    emb = model.encode([text])
    emb = np.array(emb).astype("float32")  # ✅ IMPORTANT FIX

    dim = emb.shape[1]

    key = f"{user_id}_{subject}"  # ✅ USER ISOLATION FIX

    if key not in indexes:
        indexes[key] = faiss.IndexFlatL2(dim)
        documents[key] = []

    indexes[key].add(emb)
    documents[key].append(text)


# -----------------------------
# FILE UPLOAD SUPPORT
# -----------------------------
def add_content_from_file(user_id: str, subject: str, file: UploadFile):

    content = ""

    try:
        if file.filename.endswith(".pdf"):
            with pdfplumber.open(file.file) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        content += page_text + "\n"

        elif file.filename.endswith(".docx"):
            doc = Document(file.file)
            for p in doc.paragraphs:
                content += p.text + "\n"

        elif file.filename.endswith(".txt"):
            content = file.file.read().decode("utf-8")

    except Exception as e:
        print("File parsing error:", e)
        return

    if content.strip():
        add_content(user_id, subject, content)


# -----------------------------
# RETRIEVE CONTEXT (FIXED + SAFE)
# -----------------------------
def retrieve_context(user_id: str, subject: str, k: int = 3):

    key = f"{user_id}_{subject}"

    if key not in indexes or len(documents.get(key, [])) == 0:
        return ""

    query = model.encode(["important concepts"])
    query = np.array(query).astype("float32")

    D, I = indexes[key].search(query, k)

    results = []
    for i in I[0]:
        if 0 <= i < len(documents[key]):
            results.append(documents[key][i])

    return "\n".join(results)


# -----------------------------
# GET ALL NOTES (FIXED)
# -----------------------------
def get_all_notes(user_id: str, subject: str):

    key = f"{user_id}_{subject}"

    if key not in documents:
        return ""

    return "\n\n".join(documents[key])